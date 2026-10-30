"""Builds the playtest spreadsheet (Runs, Upgrades picked, Upgrade order, Summary) from the
JSON lines tools/playtest.gd writes to <out>.jsonl. Run LibreOffice recalc afterwards so the
summary formulas have values."""
import json, glob
from collections import Counter
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
import sys, os
# Usage: python3 tools/playtest_sheet.py <folder with the playtest *.jsonl files> [out.xlsx]
S=os.path.join(sys.argv[1], "")
OUT=sys.argv[2] if len(sys.argv) > 2 else "mage_playtest_results.xlsx"
names=json.load(open(os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', 'data', 'upgrades.json')))
def flat(o):
    if isinstance(o,dict):
        if 'id' in o and 'name' in o: yield o
        for v in o.values(): yield from flat(v)
    elif isinstance(o,list):
        for v in o: yield from flat(v)
NAME={x['id']:x['name'] for x in flat(names)}
WEAP={"fire":"Fireballs","ice":"Frost Shards","thunder":"Chain Lightning","dark":"Shadow Bolts","arcane":"Arcane Bolts",
 "trunk":"Tree Trunk","wail":"Wail","silver":"Silver Bolts","spear":"Valkyrie Spear","daggers":"Twin Daggers","katana":"Katana",
 "flask":"Volatile Flasks","club":"Big Club","static":"Static Cling"}
CHAR_W={"fire_wizard":"Fireballs","ice_wizard":"Frost Shards","thunder_wizard":"Chain Lightning","dark_wizard":"Shadow Bolts","merlin":"Arcane Bolts"}
MAGE={"fire_wizard":"Fire Wizard","ice_wizard":"Ice Wizard","thunder_wizard":"Thunder Wizard","dark_wizard":"Dark Wizard","merlin":"Merlin"}
ORDER=list(MAGE)
def uname(u,ch):
    if u=="weapon_orb": return CHAR_W[ch]+" (main weapon)"
    if u.startswith("sig_"): return WEAP.get(u[4:],u)+" (Legendary)"
    return NAME.get(u,u)
runs=[]
for f in glob.glob(S+"*.jsonl"):
    for line in open(f):
        if line.strip(): runs.append(json.loads(line))
runs.sort(key=lambda r:(ORDER.index(r["char"]),r["n"]))
F="Arial"; hf=Font(name=F,bold=True,color="FFFFFF"); fill=PatternFill("solid",fgColor="3B2F5C")
thin=Side(style="thin",color="BBBBBB"); base=Font(name=F); ctr=Alignment(horizontal="center",vertical="top")
def header(ws,h,widths):
    ws.append(h)
    for c in ws[1]: c.font=hf; c.fill=fill; c.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True)
    for i,w in enumerate(widths,1): ws.column_dimensions[L(i)].width=w
    ws.freeze_panes="C2"
wb=Workbook(); ws=wb.active; ws.title="Runs"
POT=["Red Potion","Blue Potion","Purple Potion"]; PU=["Magnet","Fear Idol","5x Damage","Movement Enhancer"]
h=["Mage","Run #","Result","Time","Level","Kills","Upgrades picked (in order of first pick, × times picked)","Thog HP left"]+[p.replace(" Potion"," potions") for p in POT]+["Potions total"]+PU+["Power-ups total"]
header(ws,h,[15,6,8,7,6,8,60,9,8,8,8,9,8,8,9,10,10])
ws.row_dimensions[1].height=32
for r in runs:
    t=r["time_s"]; up=Counter(); first=[]
    for u in r["upgrades"]:
        n=uname(u,r["char"]); 
        if n not in up: first.append(n)
        up[n]+=1
    row=[MAGE[r["char"]],r["n"],r["result"].capitalize(),t/86400,r["level"],r["kills"],
         "\n".join(f"{n} ×{up[n]}" for n in first) or "-",
         (r["boss_hp"] if r["boss_hp"]>=0 else "-")]
    row+= [r["potions"].get(p,0) for p in POT]+[None]+[r["powerups"].get(p,0) for p in PU]+[None]
    ws.append(row); i=ws.max_row
    ws.cell(i,12).value=f"=SUM(I{i}:K{i})"; ws.cell(i,17).value=f"=SUM(M{i}:P{i})"
    ws.row_dimensions[i].height=max(15,15*len(first))
    for c in ws[i]: c.font=base; c.border=Border(bottom=thin); c.alignment=ctr
    ws.cell(i,1).alignment=Alignment(vertical="top"); ws.cell(i,7).alignment=Alignment(wrap_text=True,vertical="top")
    ws.cell(i,4).number_format="[m]:ss"; ws.cell(i,6).number_format="#,##0"; ws.cell(i,8).number_format="0%"
    if r["result"]=="WON":
        for c in ws[i]: c.fill=PatternFill("solid",fgColor="DFF2D8")
last=ws.max_row
ws.cell(last+2,1,"Map: Swamp (map 1). A win means killing Thog, who arrives at 8:00. Played by an automated bot that kites the horde and only grabs potions and power-ups it passes on the way, so pickup counts run lower than a human player's.").font=Font(name=F,italic=True,color="666666")
# Upgrade counts matrix
um=wb.create_sheet("Upgrades picked")
allu=[]
for r in runs:
    for u in r["upgrades"]:
        n=uname(u,r["char"]).replace(CHAR_W[r["char"]]+" (main weapon)","Main weapon")
        if n not in allu: allu.append(n)
allu.sort(key=lambda n:(n!="Main weapon","Legendary" not in n, n))
header(um,["Mage","Run #","Result"]+allu+["Total picks"],[15,6,8]+[11]*len(allu)+[9])
um.row_dimensions[1].height=45
for r in runs:
    c=Counter(uname(u,r["char"]).replace(CHAR_W[r["char"]]+" (main weapon)","Main weapon") for u in r["upgrades"])
    um.append([MAGE[r["char"]],r["n"],r["result"].capitalize()]+[c.get(n,0) or None for n in allu]+[None])
    i=um.max_row; um.cell(i,len(allu)+4).value=f"=SUM(D{i}:{L(len(allu)+3)}{i})"
    for cc in um[i]: cc.font=base; cc.alignment=Alignment(horizontal="center"); cc.border=Border(bottom=thin)
    um.cell(i,1).alignment=Alignment()
    if r["result"]=="WON":
        for cc in um[i]: cc.fill=PatternFill("solid",fgColor="DFF2D8")
um.freeze_panes="D2"
um.cell(um.max_row+2,1,"Times each upgrade was picked in each run (blank = never). Main weapon = that mage's starting weapon (Fireballs, Frost Shards, Chain Lightning, Shadow Bolts or Arcane Bolts).").font=Font(name=F,italic=True,color="666666")
# Upgrade order
uo=wb.create_sheet("Upgrade order")
header(uo,["Mage","Run #","Pick #","Upgrade"],[15,7,7,40])
for r in runs:
    for k,u in enumerate(r["upgrades"],1):
        uo.append([MAGE[r["char"]],r["n"],k,uname(u,r["char"])])
        for c in uo[uo.max_row]: c.font=base
        for j in (2,3): uo.cell(uo.max_row,j).alignment=ctr
# Summary
s=wb.create_sheet("Summary")
header(s,["Mage","Runs","Wins","Win rate","Earliest death","Latest death","Average death","Avg level","Avg kills","Avg potions","Avg power-ups"],[15,7,7,9,10,10,10,9,10,10,10])
s.row_dimensions[1].height=32
R=f"Runs!$A$2:$A${last}"; C=f"Runs!$C$2:$C${last}"; col=lambda x:f"Runs!${x}$2:${x}${last}"
for i,m in enumerate(MAGE.values(),2):
    d=f'COUNTIFS({R},A{i},{C},"Died")'
    s.append([m,f"=COUNTIF({R},A{i})",f'=COUNTIFS({R},A{i},{C},"Won")',f'=IF(B{i}=0,"-",C{i}/B{i})',
      f'=IF({d}=0,"-",_xlfn.MINIFS({col("D")},{R},A{i},{C},"Died"))',
      f'=IF({d}=0,"-",_xlfn.MAXIFS({col("D")},{R},A{i},{C},"Died"))',
      f'=IF({d}=0,"-",AVERAGEIFS({col("D")},{R},A{i},{C},"Died"))',
      f'=IF(B{i}=0,"-",AVERAGEIF({R},A{i},{col("E")}))',f'=IF(B{i}=0,"-",AVERAGEIF({R},A{i},{col("F")}))',
      f'=IF(B{i}=0,"-",AVERAGEIF({R},A{i},{col("L")}))',f'=IF(B{i}=0,"-",AVERAGEIF({R},A{i},{col("Q")}))'])
    for c in s[i]: c.font=base; c.border=Border(bottom=thin); c.alignment=Alignment(horizontal="center")
    s.cell(i,1).alignment=Alignment(); s.cell(i,4).number_format="0%"
    for j in (5,6,7): s.cell(i,j).number_format="[m]:ss"
    s.cell(i,8).number_format="0.0"; s.cell(i,9).number_format="#,##0"
    for j in (10,11): s.cell(i,j).number_format="0.0"
t=7; s.append(["Total","=SUM(B2:B6)","=SUM(C2:C6)",'=IF(B7=0,"-",C7/B7)'])
for c in s[t]: c.font=Font(name=F,bold=True); c.alignment=Alignment(horizontal="center")
s.cell(t,1).alignment=Alignment(); s.cell(t,4).number_format="0%"
s.cell(9,1,"Calculated from the Runs tab. The Upgrade order tab lists every level-up pick in sequence.").font=Font(name=F,italic=True,color="666666")
wb.save(OUT); print(len(runs),"runs")
